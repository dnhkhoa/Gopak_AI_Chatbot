from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass, field
from typing import Sequence

from openpyxl.utils import get_column_letter


@dataclass(frozen=True)
class HeaderDetection:
    header_row_index: int
    start_column_index: int
    end_column_index: int
    score: float
    columns: list[str]
    first_data_row_index: int | None = None
    last_data_row_index: int | None = None
    confidence: float = 0.0
    evidence: list[str] = field(default_factory=list)


SEMANTIC_TERMS = {
    "no",
    "may",
    "machine",
    "thoi gian bat dau",
    "start",
    "thoi gian ket thuc",
    "end",
    "thoi luong",
    "duration",
    "ten ton that",
    "loss",
    "nhom ton that",
    "loai ton that",
    "note",
}

METADATA_KEYS = {
    "export time",
    "from time",
    "to time",
    "reporter",
    "total results",
    "factory",
    "workshop",
    "line",
    "machine",
    "tu thoi gian",
    "den thoi gian",
}


def _clean(value: object) -> str:
    return "" if value is None else str(value).strip()


def _norm(value: str) -> str:
    normalized = unicodedata.normalize("NFKD", value.lower().replace("đ", "d").replace("Đ", "d"))
    return "".join(ch for ch in normalized if not unicodedata.combining(ch))


def _is_metadata_row(values: list[str]) -> bool:
    nonempty = [value for value in values if value]
    if not nonempty:
        return True
    normalized = [_norm(value).rstrip(":") for value in nonempty]
    if any(value in METADATA_KEYS or value.endswith(" time") for value in normalized):
        return True
    if len(nonempty) <= 5 and any(value.endswith(":") for value in nonempty):
        return True
    return False


def _semantic_hits(values: list[str]) -> int:
    normalized = [_norm(value) for value in values if value]
    return sum(1 for value in normalized if any(term in value for term in SEMANTIC_TERMS))


def _contiguous_segments(nonempty_positions: list[int]) -> list[tuple[int, int]]:
    if not nonempty_positions:
        return []
    segments: list[tuple[int, int]] = []
    start = prev = nonempty_positions[0]
    for pos in nonempty_positions[1:]:
        if pos == prev + 1:
            prev = pos
            continue
        segments.append((start, prev))
        start = prev = pos
    segments.append((start, prev))
    return segments


def _row_density(row: Sequence[object], start: int, end: int) -> float:
    values = [_clean(value) for value in list(row)[start : end + 1]]
    return sum(1 for value in values if value) / max(1, len(values))


def _is_repeated_header(values: list[str], headers: list[str]) -> bool:
    matches = sum(1 for a, b in zip(values, headers) if a and b and _norm(a) == _norm(b))
    return matches >= max(2, len([header for header in headers if header]) // 2)


def detect_table_bounds(
    rows: Sequence[Sequence[object]],
    header_idx: int,
    start: int,
    end: int,
    headers: list[str],
) -> tuple[int | None, int | None, list[str]]:
    evidence: list[str] = []
    first_data: int | None = None
    last_data: int | None = None
    blank_streak = 0
    for idx in range(header_idx + 1, len(rows)):
        row = list(rows[idx])
        values = [_clean(value) for value in row[start : end + 1]]
        density = sum(1 for value in values if value) / max(1, len(values))
        row_text = " ".join(_norm(value) for value in values if value)
        if _is_repeated_header(values, headers):
            evidence.append(f"ignored repeated header at row {idx + 1}")
            continue
        if row_text.startswith(("total", "grand total", "tong", "tong cong")):
            evidence.append(f"footer starts at row {idx + 1}")
            break
        if density == 0:
            blank_streak += 1
            if first_data is not None and blank_streak >= 3:
                evidence.append(f"data ended before blank streak at row {idx + 1}")
                break
            continue
        blank_streak = 0
        if density < 0.20:
            continue
        if first_data is None:
            first_data = idx
        last_data = idx
    return first_data, last_data, evidence


def detect_header(rows: Sequence[Sequence[object]], min_columns: int = 2) -> HeaderDetection | None:
    best: HeaderDetection | None = None
    scan_rows = rows[: max(200, min(len(rows), 200))]
    for idx, row in enumerate(scan_rows):
        values = [_clean(v) for v in row]
        nonempty_positions = [i for i, value in enumerate(values) if value]
        if len(nonempty_positions) < min_columns or _is_metadata_row(values):
            continue
        for start, end in _contiguous_segments(nonempty_positions):
            segment_values = values[start : end + 1]
            segment_nonempty = [value for value in segment_values if value]
            if len(segment_nonempty) < min_columns:
                continue
            unique = len({_norm(value) for value in segment_nonempty})
            text_like = sum(bool(re.search(r"[A-Za-zÀ-ỹ_]", value)) for value in segment_nonempty)
            unique_ratio = unique / max(1, len(segment_nonempty))
            semantic = _semantic_hits(segment_values)
            density_rows = [_row_density(next_row, start, end) for next_row in rows[idx + 1 : idx + 11]]
            data_density = sum(density_rows) / max(1, len(density_rows))
            first_data, last_data, bound_evidence = detect_table_bounds(rows, idx, start, end, segment_values)
            consecutive = 0 if first_data is None or last_data is None else max(0, last_data - first_data + 1)
            score = (
                len(segment_nonempty) * 1.4
                + unique_ratio * 6
                + text_like * 0.8
                + semantic * 4
                + data_density * 8
                + min(consecutive, 25) * 0.35
            )
            confidence = min(0.99, max(0.0, score / 60))
            evidence = [
                f"row={idx + 1}",
                f"columns={get_column_letter(start + 1)}:{get_column_letter(end + 1)}",
                f"semantic_hits={semantic}",
                f"data_density={data_density:.2f}",
                f"consecutive_records={consecutive}",
                *bound_evidence[:3],
            ]
            candidate = HeaderDetection(
                header_row_index=idx,
                start_column_index=start,
                end_column_index=end,
                score=score,
                columns=segment_values,
                first_data_row_index=first_data,
                last_data_row_index=last_data,
                confidence=confidence,
                evidence=evidence,
            )
            if best is None or candidate.score > best.score:
                best = candidate
    return best

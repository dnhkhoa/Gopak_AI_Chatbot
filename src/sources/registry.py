from __future__ import annotations

from dataclasses import dataclass, field
from datetime import datetime, timezone
from hashlib import sha256
import json
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from src.config import SOURCE_REGISTRY_PATH, Settings, get_settings


@dataclass(frozen=True)
class ProductionSource:
    source_id: str
    display_name: str
    workbook_path: Path
    business_domain: str
    schema_version: str
    checksum: str | None
    schema_fingerprint: str | None
    primary_grain: str
    timestamp_column: str
    timestamp_storage_timezone: str
    business_timezone: str
    supported_metrics: tuple[str, ...]
    entity_keys: tuple[str, ...]
    approved_relationships: tuple[dict[str, Any], ...] = field(default_factory=tuple)
    status: str = "unavailable"
    status_reason: str | None = None
    last_successful_ingest_time: str | None = None

    def public_dict(self) -> dict[str, Any]:
        return {
            "source_id": self.source_id,
            "display_name": self.display_name,
            "workbook_path": str(self.workbook_path),
            "business_domain": self.business_domain,
            "schema_version": self.schema_version,
            "checksum": self.checksum,
            "schema_fingerprint": self.schema_fingerprint,
            "primary_grain": self.primary_grain,
            "timestamp_column": self.timestamp_column,
            "timestamp_storage_timezone": self.timestamp_storage_timezone,
            "business_timezone": self.business_timezone,
            "supported_metrics": list(self.supported_metrics),
            "entity_keys": list(self.entity_keys),
            "approved_relationships": list(self.approved_relationships),
            "status": self.status,
            "status_reason": self.status_reason,
            "last_successful_ingest_time": self.last_successful_ingest_time,
        }


def _file_sha256(path: Path) -> str:
    digest = sha256()
    with path.open("rb") as fh:
        for block in iter(lambda: fh.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _schema_fingerprint(path: Path) -> str | None:
    if not path.exists():
        return None
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        parts: list[str] = []
        for worksheet in workbook.worksheets:
            best_row: list[str] = []
            for row in worksheet.iter_rows(min_row=1, max_row=min(60, worksheet.max_row or 1), values_only=True):
                values = [str(value).strip() for value in row if value is not None and str(value).strip()]
                if len(values) > len(best_row):
                    best_row = values
            parts.append(f"{worksheet.title}:{'|'.join(best_row)}")
        return sha256("\n".join(parts).encode("utf-8")).hexdigest()
    finally:
        workbook.close()


class SourceRegistry:
    def __init__(self, settings: Settings | None = None, registry_path: Path | None = None) -> None:
        self.settings = settings or get_settings()
        self.registry_path = registry_path or SOURCE_REGISTRY_PATH

    def load(self) -> list[ProductionSource]:
        payload = json.loads(self.registry_path.read_text(encoding="utf-8"))
        schema_version = str(payload.get("schema_version") or "unknown")
        sources: list[ProductionSource] = []
        for item in payload.get("sources", []):
            workbook_path = self.settings.root / str(item["workbook_path"])
            checksum = _file_sha256(workbook_path) if workbook_path.exists() else None
            fingerprint = _schema_fingerprint(workbook_path) if workbook_path.exists() else None
            status = "ready" if workbook_path.exists() else "unavailable"
            status_reason = None if workbook_path.exists() else "Workbook is missing from the production bundle."
            sources.append(
                ProductionSource(
                    source_id=str(item["source_id"]),
                    display_name=str(item["display_name"]),
                    workbook_path=workbook_path,
                    business_domain=str(item.get("business_domain") or ""),
                    schema_version=schema_version,
                    checksum=checksum,
                    schema_fingerprint=fingerprint,
                    primary_grain=str(item.get("primary_grain") or ""),
                    timestamp_column=str(item.get("timestamp_column") or ""),
                    timestamp_storage_timezone=str(item.get("timestamp_storage_timezone") or ""),
                    business_timezone= "Asia/Ho_Chi_Minh",
                    supported_metrics=tuple(str(value) for value in item.get("supported_metrics", [])),
                    entity_keys=tuple(str(value) for value in item.get("entity_keys", [])),
                    approved_relationships=tuple(dict(value) for value in item.get("approved_relationships", [])),
                    status=status,
                    status_reason=status_reason,
                    last_successful_ingest_time=datetime.now(timezone.utc).isoformat() if status == "ready" else None,
                )
            )
        return sources

    def by_id(self) -> dict[str, ProductionSource]:
        return {source.source_id: source for source in self.load()}

    def workbook_paths(self) -> list[Path]:
        return [source.workbook_path for source in self.load() if source.status == "ready"]

    def health(self) -> dict[str, Any]:
        sources = self.load()
        timezone_configured = bool(self.settings.business_timezone)
        return {
            "bundle_ready": bool(sources) and all(source.status == "ready" for source in sources),
            "business_timezone_configured": timezone_configured,
            "performance_formula_configured": self.settings.performance_formula_mode not in {"", "disabled"},
            "sources": [source.public_dict() for source in sources],
        }


def get_source_registry(settings: Settings | None = None) -> SourceRegistry:
    return SourceRegistry(settings)
